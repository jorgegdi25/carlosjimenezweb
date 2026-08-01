from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output" / "pdf"
DOCX_PATH = OUTPUT_DIR / "guia-practica-neurocalma-444.docx"
HERO = ROOT / "public" / "images" / "landing-neurocalma" / "image neur.png"
AUTHOR = ROOT / "public" / "assets" / "img" / "carlos-2.png"

NAVY = "082D3F"
GREEN = "28A878"
TEAL = "0E7A77"
MINT = "E9F7F1"
PALE = "F3F8F7"
GOLD = "E7B84B"
INK = "17313A"
MUTED = "557078"
WHITE = "FFFFFF"
LINE = "CFE3DD"


def rgb(value):
    return RGBColor.from_string(value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=150, bottom=120, end=150):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table, color=LINE, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:color"), color)


def remove_table_borders(table):
    set_table_borders(table, color=WHITE, size=0)


def set_col_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_font(run, size=10.5, color=INK, bold=False, italic=False, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def style_paragraph(paragraph, before=0, after=7, line=1.16, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def add_text(doc, text="", size=10.5, color=INK, bold=False, italic=False,
             align=None, before=0, after=7, line=1.16, keep=False):
    p = doc.add_paragraph()
    style_paragraph(p, before, after, line, keep)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_font(run, size, color, bold, italic)
    return p


def add_heading(doc, text, level=1):
    sizes = {1: 20, 2: 15, 3: 11.5}
    colors = {1: NAVY, 2: TEAL, 3: GREEN}
    befores = {1: 8, 2: 6, 3: 4}
    afters = {1: 9, 2: 6, 3: 4}
    p = doc.add_paragraph()
    style_paragraph(p, befores[level], afters[level], 1.05, True)
    p.style = f"Heading {level}"
    run = p.add_run(text)
    set_font(run, sizes[level], colors[level], True)
    return p


def add_rule(doc, color=GREEN, width="18"):
    p = doc.add_paragraph()
    style_paragraph(p, 1, 7, 1)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), width)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    return p


def add_bullet(doc, text, color=INK):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, 0, 4, 1.12)
    run = p.add_run(text)
    set_font(run, 10.3, color)
    return p


CURRENT_NUM_ID = None


def new_numbering_instance(doc):
    numbering = doc.part.numbering_part.element
    style_num_id = doc.styles["List Number"].element.pPr.numPr.numId.val
    source_num = numbering.num_having_numId(style_num_id)
    abstract_num_id = source_num.abstractNumId.val
    existing_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing_ids) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_num_id))
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def add_numbered(doc, title, detail):
    global CURRENT_NUM_ID
    if CURRENT_NUM_ID is None:
        CURRENT_NUM_ID = new_numbering_instance(doc)
    p = doc.add_paragraph()
    style_paragraph(p, 0, 5, 1.12)
    p.paragraph_format.left_indent = Inches(0.32)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(CURRENT_NUM_ID))
    num_pr.append(ilvl)
    num_pr.append(num_id)
    p_pr.append(num_pr)
    r1 = p.add_run(title + " ")
    set_font(r1, 10.3, NAVY, True)
    r2 = p.add_run(detail)
    set_font(r2, 10.3, INK)
    return p


def add_callout(doc, label, text, fill=MINT, accent=GREEN):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(table, [0.18, 6.25])
    remove_table_borders(table)
    set_cell_shading(table.cell(0, 0), accent)
    set_cell_shading(table.cell(0, 1), fill)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, 140, 160, 140, 160)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = table.cell(0, 1).paragraphs[0]
    style_paragraph(p, 0, 0, 1.12)
    r1 = p.add_run(label.upper() + "\n")
    set_font(r1, 9, accent, True)
    r2 = p.add_run(text)
    set_font(r2, 10.2, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_checkline(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, 0, 5, 1.1)
    r1 = p.add_run("☐  ")
    set_font(r1, 12, GREEN)
    r2 = p.add_run(text)
    set_font(r2, 10.2, INK)
    return p


def add_page_title(doc, kicker, title, subtitle=None):
    global CURRENT_NUM_ID
    CURRENT_NUM_ID = None
    add_text(doc, kicker.upper(), 9.5, GREEN, True, before=0, after=3, keep=True)
    add_heading(doc, title, 1)
    if subtitle:
        add_text(doc, subtitle, 11.2, MUTED, False, False, after=10, line=1.18)
    add_rule(doc)


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_practice_header(doc, objective, duration, when):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(table, [2.6, 1.25, 2.6])
    set_table_borders(table, LINE, 5)
    values = [
        ("PROPÓSITO", objective),
        ("DURACIÓN", duration),
        ("ÚSALA CUANDO", when),
    ]
    for idx, (label, value) in enumerate(values):
        cell = table.cell(0, idx)
        set_cell_shading(cell, PALE)
        set_cell_margins(cell, 130, 150, 130, 150)
        p = cell.paragraphs[0]
        style_paragraph(p, 0, 0, 1.05)
        r1 = p.add_run(label + "\n")
        set_font(r1, 8.2, GREEN, True)
        r2 = p.add_run(value)
        set_font(r2, 9.5, INK, idx == 1)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_log_table(doc, rows=5):
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(table, [0.7, 1.35, 1.0, 1.0, 2.45])
    set_table_borders(table, LINE, 6)
    headers = ["Día", "Práctica", "Antes\n0-10", "Después\n0-10", "¿Qué observé?"]
    for idx, text in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, 110, 120, 110, 120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, 0, 0, 1)
        run = p.add_run(text)
        set_font(run, 8.8, WHITE, True)
    set_repeat_table_header(table.rows[0])
    for _ in range(rows):
        cells = table.add_row().cells
        set_repeat_no_split(table.rows[-1])
        for idx, cell in enumerate(cells):
            set_cell_margins(cell, 180, 120, 180, 120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, 0, 0, 1)
            run = p.add_run(" ")
            set_font(run, 9.5, INK)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.16

    for style_name, size, color in [
        ("Heading 1", 20, NAVY),
        ("Heading 2", 15, TEAL),
        ("Heading 3", 11.5, GREEN),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.3)
        style.paragraph_format.left_indent = Inches(0.28)
        style.paragraph_format.first_line_indent = Inches(-0.18)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(p, 0, 0, 1)
    run = p.add_run("NEUROCALMA 444  |  GUÍA PRÁCTICA")
    set_font(run, 8.5, MUTED, True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, 0, 0, 1)
    r1 = p.add_run("Carlos Alberto Jiménez Vélez   •   ")
    set_font(r1, 8.5, MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, 4, 12, 1)
    p.add_run().add_picture(str(HERO), width=Inches(6.7))
    add_text(doc, "CURSO ONLINE · MATERIAL COMPLEMENTARIO", 9.5, GREEN, True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
    add_text(doc, "Neurocalma 444", 30, NAVY, True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=3, line=1)
    add_text(doc, "Guía práctica de ejercicios", 17, TEAL, False,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=9, line=1)
    add_text(
        doc,
        "Respiración, tapping, atención corporal y rutinas sencillas para acompañar los videos.",
        11.5, MUTED, False, False, WD_ALIGN_PARAGRAPH.CENTER, after=16, line=1.2
    )
    add_rule(doc, GOLD, "24")
    add_text(doc, "Carlos Alberto Jiménez Vélez", 11, NAVY, True,
             align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=1)
    add_text(doc, "Escritor y neuropedagogo", 9.5, MUTED,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

    # Welcome
    page_break(doc)
    add_page_title(doc, "Antes de comenzar", "Bienvenida", "Esta guía organiza las prácticas presentadas por el autor en los videos de Neurocalma 444.")
    add_text(
        doc,
        "No necesitas completar todos los ejercicios de una sola vez. Avanza a tu ritmo, elige una práctica y observa cómo responde tu cuerpo. La intención no es hacerlo perfecto, sino construir una pausa consciente que puedas repetir en la vida cotidiana."
    )
    add_callout(
        doc,
        "Idea central",
        "Primero observa. Después practica. Finalmente registra qué cambió en tu respiración, tu tensión corporal o tu estado de ánimo."
    )
    add_heading(doc, "Cómo aprovechar el curso", 2)
    add_numbered(doc, "Mira el video completo.", "La explicación y la demostración del autor son la referencia principal.")
    add_numbered(doc, "Practica con suavidad.", "Evita forzar el aire, el cuello, la mandíbula o cualquier zona sensible.")
    add_numbered(doc, "Califica tu estado.", "Antes y después, asigna un número de 0 a 10 a la intensidad del estrés o la tensión.")
    add_numbered(doc, "Repite lo que te sirve.", "Una práctica breve y constante suele ser más útil que una sesión larga y aislada.")
    add_heading(doc, "Tu punto de partida", 2)
    add_text(doc, "Hoy llego al curso con una intensidad de tensión de:  ____ / 10", 11, NAVY, True)
    add_text(doc, "La situación que más ocupa mi atención en este momento es:")
    add_rule(doc, LINE, "6")
    add_rule(doc, LINE, "6")

    # Safety
    page_break(doc)
    add_page_title(doc, "Cuidado personal", "Practica de forma segura")
    add_callout(
        doc,
        "Aviso importante",
        "Este curso y esta guía tienen fines educativos y de bienestar general. No sustituyen la evaluación, el diagnóstico ni el tratamiento de profesionales de la salud.",
        fill="FFF6DC", accent="B57B00"
    )
    add_heading(doc, "Durante las prácticas", 2)
    for item in [
        "Respira sin forzar. Si aparece mareo, hormigueo intenso, dolor o sensación de ahogo, vuelve a tu respiración natural y detén el ejercicio.",
        "Si retener el aire resulta incómodo, utiliza la respiración 4-4-4 sin pausa o simplemente alarga suavemente la exhalación.",
        "Realiza masajes y tapping con presión ligera. Evita heridas, inflamaciones o zonas dolorosas.",
        "No apliques hielo directamente sobre la piel. Si usas una compresa fresca, envuélvela y limita el contacto.",
        "Los instrumentos de vibración o masaje deben usarse según las instrucciones del fabricante y nunca sobre zonas lesionadas.",
        "Consulta a un profesional si tienes una condición respiratoria, cardiovascular, neurológica, estás en embarazo o presentas síntomas persistentes."
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Busca ayuda inmediata", 2)
    add_text(
        doc,
        "Si atraviesas una crisis emocional, tienes pensamientos de hacerte daño, dolor en el pecho, dificultad respiratoria intensa u otra emergencia, busca atención profesional o los servicios de emergencia de tu localidad."
    )

    # Course route
    page_break(doc)
    add_page_title(doc, "Ruta de aprendizaje", "Mapa del curso")
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(table, [0.9, 2.2, 3.35])
    set_table_borders(table, LINE, 6)
    for idx, h in enumerate(["VIDEO", "TEMA", "PRÁCTICA PRINCIPAL"]):
        c = table.cell(0, idx)
        set_cell_shading(c, NAVY)
        set_cell_margins(c)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, 0, 0, 1)
        set_font(p.add_run(h), 9, WHITE, True)
    rows = [
        ("Intro", "Presentación", "Respiración y reconocimiento corporal"),
        ("1", "Nervio vago", "Masaje, tarareo y contacto corporal"),
        ("2", "Tapping", "Secuencia y escala subjetiva 0-10"),
        ("3", "Respuesta de estrés", "Comprender alerta y calma"),
        ("4", "Respiración", "Ritmos 4-4-4 y 4-7-8"),
        ("5", "Práctica integrada", "Respiración, gratitud, tapping y masaje"),
        ("6", "Oídos y visualización", "Sonido, masaje y observación sin juicio"),
        ("7", "Motivación", "Reflexión sobre hábitos y recompensa"),
        ("8", "Alimentación", "Observación de hábitos cotidianos"),
    ]
    for video, topic, practice in rows:
        cells = table.add_row().cells
        set_repeat_no_split(table.rows[-1])
        for idx, value in enumerate((video, topic, practice)):
            set_cell_margins(cells[idx], 105, 130, 105, 130)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, 0, 0, 1.05)
            set_font(p.add_run(value), 9.2, NAVY if idx == 0 else INK, idx == 0)

    # Breathing 444
    page_break(doc)
    add_page_title(doc, "Práctica 1", "Respiración 4-4-4", "Una pausa rítmica para dirigir la atención al cuerpo.")
    add_practice_header(doc, "Regular el ritmo y volver al presente", "2-4 min", "Necesites una pausa sencilla")
    add_heading(doc, "Preparación", 2)
    add_bullet(doc, "Siéntate con la espalda cómoda y los pies apoyados.")
    add_bullet(doc, "Afloja hombros, mandíbula y manos.")
    add_bullet(doc, "Respira por la nariz si te resulta cómodo.")
    add_heading(doc, "Paso a paso", 2)
    add_numbered(doc, "Inhala durante 4 tiempos.", "Permite que el abdomen se expanda sin levantar demasiado los hombros.")
    add_numbered(doc, "Haz una pausa de 4 tiempos.", "Solo si se siente natural; si no, continúa sin retención.")
    add_numbered(doc, "Exhala durante 4 tiempos.", "Suelta el aire suavemente, sin vaciarte a la fuerza.")
    add_numbered(doc, "Repite entre 4 y 8 ciclos.", "Mantén un ritmo cómodo y estable.")
    add_callout(doc, "Frase de apoyo", "Al inhalar: «recibo calma». Al exhalar: «suelto tensión».")
    add_heading(doc, "Registro rápido", 2)
    add_text(doc, "Antes: ____ / 10      Después: ____ / 10")
    add_text(doc, "Noté cambios en:  ☐ respiración   ☐ hombros   ☐ mandíbula   ☐ pensamientos   ☐ otro: __________")

    # Breathing 478
    page_break(doc)
    add_page_title(doc, "Práctica 2", "Respiración 4-7-8", "Una variante del curso para practicar con calma y sin exigencia.")
    add_practice_header(doc, "Alargar la exhalación y enfocar la atención", "2-3 min", "Ya domines un ritmo suave")
    add_heading(doc, "Paso a paso", 2)
    add_numbered(doc, "Inhala durante 4 tiempos.", "Respira de forma silenciosa y cómoda.")
    add_numbered(doc, "Pausa durante 7 tiempos.", "Omite o acorta esta parte si te genera incomodidad.")
    add_numbered(doc, "Exhala durante 8 tiempos.", "Deja salir el aire lentamente, sin empujar.")
    add_numbered(doc, "Realiza de 2 a 4 ciclos.", "Después vuelve a respirar con naturalidad.")
    add_callout(
        doc,
        "Alternativa suave",
        "Si 4-7-8 resulta exigente, utiliza 4-4-6: inhala 4, pausa 4 y exhala 6. La comodidad es más importante que cumplir una cifra."
    )
    add_heading(doc, "Chequeo corporal", 2)
    for item in [
        "¿Mi respiración se siente más amplia o más tranquila?",
        "¿Puedo soltar un poco los hombros?",
        "¿Necesito continuar o es suficiente por ahora?"
    ]:
        add_checkline(doc, item)

    # Tapping
    page_break(doc)
    add_page_title(doc, "Práctica 3", "Secuencia de tapping", "Acompaña esta ficha con la demostración del capítulo 2.")
    add_practice_header(doc, "Observar una emoción y reducir su intensidad subjetiva", "5-8 min", "Una situación sigue dando vueltas")
    add_heading(doc, "Antes de empezar", 2)
    add_numbered(doc, "Elige una situación concreta.", "Trabaja con algo manejable, no con el recuerdo más difícil de tu vida.")
    add_numbered(doc, "Ponle una intensidad de 0 a 10.", "Cero significa sin tensión; diez, tensión muy alta.")
    add_numbered(doc, "Formula una frase.", "Ejemplo: «Aunque siento esta tensión, me observo con respeto y acepto avanzar poco a poco».")
    add_heading(doc, "Recorrido orientativo", 2)
    table = doc.add_table(rows=3, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(table, [2.15, 2.15, 2.15])
    remove_table_borders(table)
    points = [
        ("1", "Borde de la mano"),
        ("2", "Inicio de la ceja"),
        ("3", "Lado del ojo"),
        ("4", "Debajo del ojo"),
        ("5", "Debajo de la nariz"),
        ("6", "Mentón"),
        ("7", "Zona de clavículas"),
        ("8", "Parte superior de la cabeza"),
        ("9", "Manos y dedos, según el video"),
    ]
    for idx, (num, label) in enumerate(points):
        cell = table.cell(idx // 3, idx % 3)
        set_cell_shading(cell, MINT if idx % 2 == 0 else PALE)
        set_cell_margins(cell, 130, 140, 130, 140)
        p = cell.paragraphs[0]
        style_paragraph(p, 0, 0, 1.05)
        r1 = p.add_run(num + "  ")
        set_font(r1, 11, GREEN, True)
        set_font(p.add_run(label), 9.4, INK)
    add_text(doc, "Da golpecitos suaves y sigue el orden, la cantidad y las zonas mostradas por el autor.", 9.3, MUTED, False, True, after=8)
    add_heading(doc, "Al terminar", 2)
    add_text(doc, "Intensidad inicial: ____ / 10      Intensidad final: ____ / 10")
    add_text(doc, "La emoción cambió de forma, lugar o intensidad de esta manera:")
    add_rule(doc, LINE, "6")

    # Massage
    page_break(doc)
    add_page_title(doc, "Práctica 4", "Masaje, sonido y atención corporal", "Movimientos suaves inspirados en los capítulos 1, 5 y 6.")
    add_practice_header(doc, "Liberar tensión y dirigir la atención al cuerpo", "5-7 min", "Sientas rigidez o agitación")
    add_heading(doc, "Recorrido sugerido", 2)
    steps = [
        ("Rostro", "Desliza suavemente los dedos por frente, cejas y pómulos."),
        ("Mandíbula", "Abre y cierra la boca con suavidad; masajea sin presionar la articulación."),
        ("Orejas", "Recorre el borde externo y masajea alrededor de la oreja con movimientos pequeños."),
        ("Cuello", "Toca los costados con mucha suavidad; evita presión directa o fuerte."),
        ("Pecho y abdomen", "Apoya las manos, percibe el movimiento respiratorio y no fuerces la zona."),
        ("Sonido", "Tararea de forma cómoda durante una exhalación y observa la vibración."),
    ]
    for idx, (title, detail) in enumerate(steps, 1):
        add_numbered(doc, f"{title}.", detail)
    add_callout(doc, "Cierre", "Coloca una mano en el pecho y otra en el abdomen. Haz tres respiraciones naturales y nota qué zona se siente diferente.")
    add_heading(doc, "Mi observación", 2)
    add_text(doc, "La zona que acumulaba más tensión era: __________________________________")
    add_text(doc, "Después de la práctica se siente: ________________________________________")

    # Visualization
    page_break(doc)
    add_page_title(doc, "Práctica 5", "Visualización y observación sin juicio", "Una manera de mirar la experiencia sin quedar atrapado en ella.")
    add_practice_header(doc, "Crear distancia y reconocer sensaciones", "4-6 min", "Necesites ordenar pensamientos")
    add_heading(doc, "Paso a paso", 2)
    add_numbered(doc, "Elige una situación reciente.", "Procura que sea una experiencia que puedas observar sin desbordarte.")
    add_numbered(doc, "Nombra lo que ocurre.", "«Estoy notando preocupación», «siento presión en el pecho» o «mi mente está acelerada».")
    add_numbered(doc, "Ubica la sensación.", "¿Está en la garganta, el abdomen, la mandíbula, la frente o en otro lugar?")
    add_numbered(doc, "Respira y observa.", "No intentes borrar la sensación; mira si cambia mientras respiras.")
    add_numbered(doc, "Cierra con una acción posible.", "Elige un paso pequeño que puedas realizar hoy.")
    add_callout(doc, "Recuerda", "Observar sin juicio no significa aprobar lo ocurrido. Significa reconocer la experiencia antes de decidir cómo responder.")
    add_heading(doc, "Escritura breve", 2)
    prompts = [
        "La situación que observo es:",
        "En mi cuerpo aparece como:",
        "El paso pequeño que puedo dar es:",
    ]
    for prompt in prompts:
        add_text(doc, prompt, 10.2, NAVY, True, after=2)
        add_rule(doc, LINE, "6")

    # 5 min routine
    page_break(doc)
    add_page_title(doc, "Rutina diaria", "Neurocalma en 5 minutos", "Una versión breve para repetir al comenzar el día o durante una pausa.")
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(table, [0.8, 2.1, 3.55])
    set_table_borders(table, LINE, 6)
    for idx, value in enumerate(("TIEMPO", "ACCIÓN", "ENFOQUE")):
        c = table.cell(0, idx)
        set_cell_shading(c, NAVY)
        set_cell_margins(c)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, 0, 0, 1)
        set_font(p.add_run(value), 9, WHITE, True)
    routine = [
        ("1 min", "Llegar", "Apoya los pies, afloja hombros y observa tu respiración."),
        ("2 min", "Respirar", "Realiza 4-4-4 a un ritmo cómodo."),
        ("1 min", "Tocar", "Masajea rostro, orejas o mandíbula con suavidad."),
        ("1 min", "Cerrar", "Tararea, agradece algo concreto y define una intención."),
    ]
    for time, action, focus in routine:
        cells = table.add_row().cells
        for idx, value in enumerate((time, action, focus)):
            set_cell_margins(cells[idx], 170, 140, 170, 140)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, 0, 0, 1.08)
            set_font(p.add_run(value), 9.6, NAVY if idx < 2 else INK, idx < 2)
    add_heading(doc, "Antes y después", 2)
    add_text(doc, "Antes: ____ / 10      Después: ____ / 10")
    add_checkline(doc, "Puedo continuar mi día con un poco más de claridad.")
    add_checkline(doc, "Necesito descansar, pedir apoyo o elegir una práctica más lenta.")

    # 15 min routine
    page_break(doc)
    add_page_title(doc, "Rutina completa", "Neurocalma en 15 minutos", "Integra las principales prácticas del curso sin apresurarte.")
    items = [
        ("Minutos 0-2", "Llegada", "Postura cómoda, escala 0-10 y tres respiraciones naturales."),
        ("Minutos 2-6", "Respiración", "Practica 4-4-4 o una variante cómoda."),
        ("Minutos 6-10", "Tapping", "Recorre la secuencia del capítulo 2 con una frase de autoaceptación."),
        ("Minutos 10-13", "Masaje y sonido", "Rostro, mandíbula, orejas y tarareo suave."),
        ("Minutos 13-15", "Visualización y cierre", "Observa la situación, agradece algo concreto y vuelve a calificarte."),
    ]
    for title, action, detail in items:
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_col_widths(table, [1.45, 5.0])
        set_table_borders(table, LINE, 5)
        set_cell_shading(table.cell(0, 0), GREEN)
        set_cell_shading(table.cell(0, 1), PALE)
        for cell in table.rows[0].cells:
            set_cell_margins(cell, 135, 150, 135, 150)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p1 = table.cell(0, 0).paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p1, 0, 0, 1)
        set_font(p1.add_run(title), 9.5, WHITE, True)
        p2 = table.cell(0, 1).paragraphs[0]
        style_paragraph(p2, 0, 0, 1.08)
        r = p2.add_run(action + ": ")
        set_font(r, 9.8, NAVY, True)
        set_font(p2.add_run(detail), 9.7, INK)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    add_callout(doc, "Frecuencia sugerida", "Prueba esta rutina dos o tres veces durante la primera semana. Ajusta la duración a tu energía y comodidad.")

    # 7-day plan
    page_break(doc)
    add_page_title(doc, "Plan práctico", "7 días de Neurocalma", "Una propuesta sencilla para convertir los videos en experiencia.")
    plan = [
        ("Día 1", "Introducción + capítulo 3", "Reconoce cómo se manifiesta el estrés en tu cuerpo."),
        ("Día 2", "Capítulo 4", "Practica respiración 4-4-4 durante 3 minutos."),
        ("Día 3", "Capítulo 1", "Realiza masaje facial, de orejas y tarareo suave."),
        ("Día 4", "Capítulo 2", "Aplica tapping a una situación manejable y registra 0-10."),
        ("Día 5", "Capítulos 5 y 6", "Completa una rutina integrada de 10 a 15 minutos."),
        ("Día 6", "Capítulos 7 y 8", "Observa hábitos de motivación, descanso y alimentación sin juzgarte."),
        ("Día 7", "Repaso", "Elige tus dos prácticas más útiles y diseña tu rutina personal."),
    ]
    for day, video, action in plan:
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_col_widths(table, [0.8, 2.1, 3.55])
        set_table_borders(table, LINE, 5)
        for idx, value in enumerate((day, video, action)):
            set_cell_margins(table.cell(0, idx), 120, 135, 120, 135)
            table.cell(0, idx).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(table.cell(0, idx), MINT if idx != 0 else NAVY)
            p = table.cell(0, idx).paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, 0, 0, 1.05)
            set_font(p.add_run(value), 9.2, WHITE if idx == 0 else INK, idx < 2)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # Log
    page_break(doc)
    add_page_title(doc, "Cuaderno de práctica", "Registro de 7 días", "Usa una fila después de cada sesión; no necesitas escribir mucho.")
    add_log_table(doc, 7)
    add_heading(doc, "Al finalizar la semana", 2)
    add_text(doc, "La práctica que más me ayudó fue: ______________________________________")
    add_text(doc, "La señal corporal que aprendí a reconocer fue: _____________________________")
    add_text(doc, "Mi rutina personal tendrá estos tres elementos:")
    add_checkline(doc, "____________________________________________________________")
    add_checkline(doc, "____________________________________________________________")
    add_checkline(doc, "____________________________________________________________")

    # Habit reflection
    page_break(doc)
    add_page_title(doc, "Reflexión complementaria", "Hábitos que acompañan la calma", "El curso invita a mirar el bienestar de forma integral.")
    add_text(
        doc,
        "Esta página no propone una dieta ni reemplaza asesoría profesional. Su propósito es ayudarte a observar hábitos cotidianos y escoger cambios realistas."
    )
    sections = [
        ("Descanso", ["Mantengo una hora aproximada para dormir.", "Reduzco pantallas o estímulos antes de acostarme."]),
        ("Movimiento", ["Incluyo caminatas, estiramientos o actividad que disfruto.", "Evito pasar periodos demasiado largos sin moverme."]),
        ("Alimentación consciente", ["Observo mi consumo de bebidas azucaradas y ultraprocesados.", "Priorizo agua y alimentos variados que tolero bien."]),
        ("Motivación y recompensa", ["Divido las tareas grandes en pasos pequeños.", "Reconozco avances sin depender de la perfección."]),
    ]
    for title, checks in sections:
        add_heading(doc, title, 3)
        for check in checks:
            add_checkline(doc, check)
    add_callout(
        doc,
        "Decisión de esta semana",
        "El hábito pequeño y específico que quiero probar es: ________________________________________________"
    )

    # Personal plan and closing
    page_break(doc)
    add_page_title(doc, "Cierre", "Mi plan personal de Neurocalma", "Elige una rutina que realmente puedas sostener.")
    prompts = [
        ("Mi señal de alerta", "Cuando noto ____________________________________, sé que necesito hacer una pausa."),
        ("Mi práctica principal", "La técnica que elegiré primero es ______________________________________."),
        ("Mi momento", "La practicaré a las __________ o después de ____________________________."),
        ("Mi apoyo", "Si necesito ayuda, hablaré con _________________________________________."),
    ]
    for title, prompt in prompts:
        add_heading(doc, title, 3)
        add_text(doc, prompt, 10.5, INK, after=7)
        add_rule(doc, LINE, "6")
    add_callout(
        doc,
        "Compromiso amable",
        "Durante los próximos 7 días practicaré sin exigirme resultados perfectos. Mi meta será observar, respirar y responder con mayor conciencia."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, 12, 5, 1)
    p.add_run().add_picture(str(AUTHOR), width=Inches(0.95))
    add_text(doc, "Carlos Alberto Jiménez Vélez", 11, NAVY, True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    add_text(doc, "Gracias por hacer de la calma una práctica cotidiana.", 10, MUTED,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

    doc.core_properties.title = "Neurocalma 444 - Guía práctica de ejercicios"
    doc.core_properties.subject = "Material complementario del curso Neurocalma 444"
    doc.core_properties.author = "Carlos Alberto Jiménez Vélez"
    doc.core_properties.keywords = "Neurocalma, respiración, tapping, bienestar, guía práctica"
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_document()
